"""parsed.json → 플레이어 시트의 '어빌리티'/'아이템' 탭과 같은 열 구성의 xlsx 생성."""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

src = sys.argv[1] if len(sys.argv) > 1 else 'parsed.json'
out = sys.argv[2] if len(sys.argv) > 2 else '블러드문_기본룰북_데이터.xlsx'
d = json.load(open(src, encoding='utf-8'))
# 서플리먼트(길티 위치즈) — 홈페이지 룰 데이터 넣기는 종류째 갈아 끼우므로 한 탭에 같이 둔다
BOOK, GW_BOOK = '기본룰북', '길티 위치즈'
gw = json.load(open('gw.json', encoding='utf-8')) if os.path.exists('gw.json') else {'abilities': [], 'items': []}


def num(v):
    """숫자 코스트는 숫자로, '-'/'없음'은 그대로."""
    return int(v) if v.isdigit() else v


def sheet(ws, header, rows, widths, fill, center_cols):
    ws.append(header)
    for r in rows:
        ws.append(r)
    head_fill = PatternFill('solid', fgColor=fill)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(horizontal='center', vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row:
            center = c.column_letter in center_cols
            c.alignment = Alignment(wrap_text=True, vertical='center',
                                    horizontal='center' if center else 'left')
            if c.column_letter == 'A':
                c.font = Font(bold=True)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()
ws = wb.active
ws.title = '어빌리티(룰북)'
sheet(ws, ['이름', '그룹', '타입', '코스트', '지정특기', '효과', '개요', '레벨제한', '비고', '룰북'],
      [[a['name'], a['group'], a['type'], num(a['cost']), a['spec'], a['effect'],
        a['flavor'], a['level'], a['note'], book]
       for book, src_ in ((BOOK, d), (GW_BOOK, gw)) for a in src_['abilities']],
      {'A': 22.8, 'B': 13, 'C': 9.6, 'D': 8, 'E': 12.9, 'F': 53.6, 'G': 40, 'H': 9.6, 'I': 19.1, 'J': 11},
      'D0E0E3', 'ABCDEH')

CAT = {'아이템': '일반', '레어아이템': '레어'}
rows = []
for n, it in enumerate(d['items'], 1):
    cat = '무장' if it['effect'].startswith('장비하고 있으면 「무장') else CAT[it['cat']]
    rows.append([n, it['name'], cat, it['effect'], it['flavor'], it['note'], BOOK])
for it in gw['items']:
    rows.append([len(rows) + 1, it['name'], it['cat'], it['effect'], it['flavor'], it['note'], GW_BOOK])
sheet(wb.create_sheet('아이템(룰북)'), ['no', '아이템명', '분류', '효과', '설명', '비고', '룰북'], rows,
      {'A': 6.5, 'B': 17.1, 'C': 8.9, 'D': 48.2, 'E': 40, 'F': 26.6, 'G': 11}, 'FCE5CD', 'ABC')

wb.save(out)
print(out, len(d['abilities']) + len(gw['abilities']), len(rows))

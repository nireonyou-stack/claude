"""기존 플레이어 시트(xlsx 내보내기)에서 길티 위치즈 서플리먼트 데이터만 gw.json으로 뽑는다.

룰 데이터 넣기는 그 종류를 통째로 갈아 끼우므로, 기본룰북 데이터와 한 탭에 같이 둬야 사라지지 않는다.
"""
import json, sys
import openpyxl

GW = {'종족:마법소녀', '조직:사립심연여학교', '조직:생츄어리', '조직:미려파', '조직:이노우에 가',
      '무장:횃불', '무장:노래', '무장:주사기', '무장:소문', '무장:고문도구', '무장:화장품', '무장:요리'}


def s(v):
    """셀 값 → 문자열. 2.0 같은 코스트는 정수로."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


wb = openpyxl.load_workbook(sys.argv[1])
ab = [dict(name=s(r[0]), group=s(r[1]), type=s(r[2]), cost=s(r[3]), spec=s(r[4]), effect=s(r[5]),
           flavor=s(r[6]), level=s(r[7]), note=s(r[8]))
      for r in wb['어빌리티'].iter_rows(min_row=2, values_only=True) if r[0] and r[1] in GW]
it = [dict(name=s(r[1]), cat=s(r[2]), effect=s(r[3]), flavor=s(r[4]), note=s(r[5]))
      for r in wb['아이템'].iter_rows(min_row=2, values_only=True) if r[5] and '길티 위치즈' in str(r[5])]
json.dump(dict(abilities=ab, items=it), open('gw.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(len(ab), len(it))

"""인세인 어빌리티 데이터 시트(xlsx 내보내기) → 홈페이지 룰 데이터 넣기(ruledata.php)용 탭 한 장.

- 넣기는 그 종류를 통째로 갈아 끼우므로 기본 · 월드 · 원무현한 어빌리티를 한 탭에 모은다.
- 타입은 구획 줄(공격 · 서포트 · 장비 어빌리티)이나 이름 뒤 '(서포트)' 에서 가져온다.
- 효과 칸의 '해설 :' 뒤는 해설 열로 나눈다(틀의 eff2).
- 구획 줄(공적점이 필요한 어빌리티 · 월드 이름)은 그대로 둔다 — 넣기가 묶음 이름(_grp)으로 쓴다.
"""
import re, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

src = sys.argv[1] if len(sys.argv) > 1 else 'src.xlsx'
out = sys.argv[2] if len(sys.argv) > 2 else '인세인_어빌리티_홈페이지.xlsx'
TYPES = ('공격', '서포트', '장비')


def txt(v):
    """셀 값 정리 — roll20 굴림 [[1D6]] 을 1D6 으로, 줄바꿈은 \n 으로."""
    if v is None:
        return ''
    s = str(v).replace('\r\n', '\n').replace('\r', '\n')
    return re.sub(r'\[\[([^\[\]]+)\]\]', r'\1', s).strip()


def skill(v):
    """지정특기 — 《》 를 떼고 한 줄로."""
    return re.sub(r'\s*\n\s*', ' ', txt(v).replace('《', '').replace('》', ''))


def split_note(eff):
    """'효과\n\n해설 : …' → (효과, 해설)."""
    m = re.split(r'\n\s*해설\s*:\s*', eff, maxsplit=1)
    return (m[0].strip(), m[1].strip()) if len(m) == 2 else (eff, '')


wb = openpyxl.load_workbook(src)
rows = []                                   # [이름, 타입, 지정특기, 효과, 해설, 습득조건, 룰북] · 구획 줄은 [이름]
base = ''
for r in list(wb['어빌리티'].iter_rows(values_only=True))[1:]:
    name, sk, eff, cond = (list(r) + [None] * 4)[:4]
    name = txt(name)
    if not name:
        continue
    if not eff:                             # 구획 줄
        head = name.replace(' 어빌리티', '')
        if head in TYPES:
            base = head
        rows.append([name])
        continue
    m = re.match(r'^(.*?)\s*\n?\((공격|서포트|장비)\)$', name, re.S)   # 월드 어빌리티 '동정/처녀\n(서포트)'
    typ = base
    if m:
        name, typ = m.group(1).strip(), m.group(2)
    e, note = split_note(txt(eff))
    c = txt(cond)
    rows.append([name, typ, skill(sk), e, note, '' if c == '-' else c, '기본룰북'])

# 원무현한 — 머리줄이 둘째 줄, B 열은 구획(합친 칸), G 열이 타입
if '원무현한' in wb.sheetnames:
    rows.append(['원무현한'])
    for r in list(wb['원무현한'].iter_rows(values_only=True))[2:]:
        _, _, name, sk, eff, note, typ = (list(r) + [None] * 7)[:7]
        if not txt(name):
            continue
        rows.append([txt(name), txt(typ), skill(sk), txt(eff), txt(note), '', '원무현한'])

wo = openpyxl.Workbook()
ws = wo.active
ws.title = '어빌리티(홈페이지)'
ws.append(['이름', '타입', '지정특기', '효과', '해설', '습득조건', '룰북'])
for r in rows:
    ws.append(r)
fill = PatternFill('solid', fgColor='D9D9D9')
for c in ws[1]:
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal='center', vertical='center')
for row in ws.iter_rows(min_row=2):
    sec = row[1].value is None              # 구획 줄은 회색으로
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='center',
                                horizontal='center' if c.column_letter in 'BCG' else 'left')
        if sec:
            c.fill = fill
            c.font = Font(bold=True)
for col, w in {'A': 16, 'B': 8, 'C': 18, 'D': 60, 'E': 36, 'F': 24, 'G': 10}.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
wo.save(out)
n = sum(1 for r in rows if len(r) > 1)
print(out, n, 'abilities', len(rows) - n, 'section rows')
